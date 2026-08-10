from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "innovation_paper_v2" / "MANUSCRIPT_POLISHED_V2.md"
OUT_DOCX = ROOT / "output" / "documents" / "SOH_innovation_paper_Energy_style_v2.docx"
FIG_DIR = ROOT / "output" / "figures" / "paper_v2"

TITLE = "小样本电池SOH估计中物理趋势约束与融合域对齐的泛化边界"
EN_TITLE = "Generalization limits of physics-guided trend constraints and fusion-domain alignment for small-sample battery SOH estimation"
BLUE = RGBColor(0x00, 0x78, 0xB5)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_run_font(run, size=8.7, bold=False, italic=False, color=None, east_asia="SimSun"):
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def configure_section(section, columns=2):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "360")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 7, color=MUTED)


def set_cell_margins(cell, value=65):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge in ("top", "start", "bottom", "end"):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(cell, fill="E9E9E9"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def borders(table, color="777777", size="4"):
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tbl_pr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), size)
        line.set(qn("w:color"), color)
        node.append(line)


def keep_row_together(row, keep_with_next=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = keep_with_next


def clean(text):
    return text.replace("**", "").replace("`", "").strip()


def add_inline(paragraph, text, size=8.6, first_indent=True):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.0
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.35)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        set_run_font(run, size=size, bold=bold)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(clean(text))
    set_run_font(run, size=10 if level == 1 else 9, bold=True, italic=(level == 2), east_asia="SimHei" if level == 1 else "SimSun")


def add_table(doc, raw_rows):
    rows = [[clean(c.strip()) for c in row.strip().strip("|").split("|")] for row in raw_rows]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    cols = max(map(len, rows))
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, values in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Inches(3.10 / cols)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 45)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(values[c_idx] if c_idx < len(values) else "")
            set_run_font(run, size=5.4 if cols >= 6 else 6.0, bold=(r_idx == 0))
            if r_idx == 0:
                shade(cell)
        keep_row_together(table.rows[r_idx], keep_with_next=(r_idx < len(rows) - 1))
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_full_figure(doc, stem, caption):
    wide = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(wide, columns=1)
    wide.header.is_linked_to_previous = True
    wide.footer.is_linked_to_previous = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(FIG_DIR / f"{stem}.png"), width=Inches(6.85))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cp.paragraph_format.keep_together = True
    cp.paragraph_format.space_after = Pt(5)
    r = cp.add_run(clean(caption))
    set_run_font(r, size=7.1)
    narrow = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(narrow, columns=2)
    narrow.header.is_linked_to_previous = True
    narrow.footer.is_linked_to_previous = True


def block(md, start, end):
    s = md.index(start) + len(start)
    e = md.index(end, s)
    return md[s:e].strip()


def build():
    md = SOURCE.read_text(encoding="utf-8")
    zh_abs = block(md, "## 摘要", "## Abstract")
    en_abs = block(md, "## Abstract", "## 1. 引言")
    zh_abs = re.sub(r"\n\n\*\*关键词：\*\*.*", "", zh_abs, flags=re.S).strip()
    en_abs = re.sub(r"\n\n\*\*Keywords:\*\*.*", "", en_abs, flags=re.S).strip()
    body = md[md.index("## 1. 引言"):]

    doc = Document()
    first = doc.sections[0]
    configure_section(first, columns=1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(8.7)
    normal.paragraph_format.space_after = Pt(2.5)

    header = first.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("SOH estimation | Generalization boundary analysis"), 7, italic=True, color=MUTED)
    add_page_field(first.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Energy-style research manuscript | Polished v2"), 8, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run(TITLE), 17, bold=True, east_asia="SimHei")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run(EN_TITLE), 10.5, italic=True, color=MUTED)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Anonymous author(s) | affiliations withheld for double-blind review"), 9, bold=True)

    info = doc.add_table(rows=2, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Inches(1.75)
    info.columns[1].width = Inches(5.05)
    for row in info.rows:
        row.cells[0].width = Inches(1.75)
        row.cells[1].width = Inches(5.05)
    for idx, label in enumerate(("A R T I C L E   I N F O", "A B S T R A C T")):
        set_run_font(info.cell(0, idx).paragraphs[0].add_run(label), 8.3, bold=True)
    left = info.cell(1, 0)
    left.text = ""
    set_run_font(left.paragraphs[0].add_run("Keywords:\nLithium-ion battery\nState of health\nSmall-sample learning\nPhysical consistency\nDomain shift\nNegative transfer"), 7.6, italic=True)
    right = info.cell(1, 1)
    right.text = ""
    add_inline(right.paragraphs[0], en_abs, size=7.4, first_indent=False)
    for row in info.rows:
        for cell in row.cells:
            set_cell_margins(cell, 80)

    p = doc.add_paragraph()
    set_run_font(p.add_run("中文摘要"), 9, bold=True)
    p = doc.add_paragraph()
    add_inline(p, zh_abs, size=8.0, first_indent=False)

    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(body_sec, columns=2)
    body_sec.header.is_linked_to_previous = True
    body_sec.footer.is_linked_to_previous = True

    lines = body.splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        fig_match = re.fullmatch(r"<!-- FIGURE:([^|]+)\|(.*) -->", stripped)
        if fig_match:
            add_full_figure(doc, fig_match.group(1), fig_match.group(2))
            i += 1
            continue
        if stripped.startswith("<!--"):
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            add_table(doc, rows)
            continue
        if stripped.startswith("## "):
            add_heading(doc, stripped[3:], 1)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 2)
            i += 1
            continue
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("<!--"):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        if stripped.startswith("**表"):
            p.paragraph_format.keep_with_next = True
        add_inline(p, " ".join(paragraph_lines), size=8.5)

    ref_count = 0
    for p in doc.paragraphs:
        if re.match(r"^\[\d+\]", p.text.strip()):
            ref_count += 1
            if ref_count == 9:
                p.insert_paragraph_before().add_run().add_break(WD_BREAK.COLUMN)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                set_run_font(run, 6.9, color=BLUE)

    props = doc.core_properties
    props.title = TITLE
    props.subject = "Auditable G-MTC reproduction, Trend-only physical constraints, and fusion-domain alignment"
    props.author = "Anonymous author(s)"
    props.keywords = "SOH, G-MTC, TimesNet, iTransformer, physical consistency, domain alignment"
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
