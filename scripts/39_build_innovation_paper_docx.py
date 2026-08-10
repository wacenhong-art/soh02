from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "innovation_paper_v1" / "MANUSCRIPT_DRAFT_V1.md"
OUT_DOCX = ROOT / "output" / "documents" / "SOH_innovation_paper_Energy_style_v1.docx"

TITLE = "从物理Trend约束到融合域对齐：小样本锂离子电池SOH估计的可信复现、负迁移与泛化边界"
EN_TITLE = "From Physics-Guided Trend Constraints to Fusion-Domain Alignment: Auditable Reproduction, Negative Transfer, and Generalization Boundaries in Small-Sample Battery SOH Estimation"

BLUE = RGBColor(0x00, 0x84, 0xC8)
MUTED = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = "E9E9E9"
GRID = "808080"


def set_run_font(run, size=9, bold=False, italic=False, color=None, east_asia="SimSun"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def set_columns(section, count: int, space_twips: int = 360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_twips))


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=7.5, color=MUTED)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def clean_md(text: str) -> str:
    text = re.sub(r"<!--.*?-->", "", text)
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = text.replace("*", "")
    return text.strip()


def add_inline(paragraph, text: str, size=9, justify=True):
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    # Preserve bold fragments from Markdown.
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if is_bold else part
        value = value.replace("`", "").replace("*", "")
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=is_bold)


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(clean_md(text))
    set_run_font(run, size=10.2 if level == 1 else 9.3, bold=True, italic=(level >= 2))
    return p


def add_markdown_table(doc, rows):
    if len(rows) < 2:
        return
    parsed = [[clean_md(c.strip()) for c in row.strip().strip("|").split("|")] for row in rows]
    if len(parsed) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in parsed[1]):
        parsed.pop(1)
    cols = max(len(r) for r in parsed)
    compact_headers = {
        "RMSE均值±标准差/%": "RMSE mean±SD/%",
        "RMSE中位数/%": "RMSE median/%",
        "MAE均值/%": "MAE mean/%",
        "MAPE均值/%": "MAPE mean/%",
        "Trend违规率均值": "Trend violation",
        "Original方法中位RMSE/%": "Original RMSE/%",
        "CILA方法中位RMSE/%": "CILA RMSE/%",
        "配对中位差/%": "ΔRMSE/%",
        "95%时间块bootstrap区间": "95% CI",
    }
    parsed[0] = [compact_headers.get(c, c) for c in parsed[0]]
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = 3.15 / cols
    for ridx, values in enumerate(parsed):
        for cidx in range(cols):
            cell = table.cell(ridx, cidx)
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(values[cidx] if cidx < len(values) else "")
            set_run_font(run, size=6.5, bold=(ridx == 0))
            if ridx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_figure(doc, filename: str, caption: str):
    path = ROOT / "output" / "figures" / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(3.05))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(5)
    run = cp.add_run(caption)
    set_run_font(run, size=7.2)


def extract_block(md: str, start: str, end: str) -> str:
    s = md.index(start) + len(start)
    e = md.index(end, s)
    return md[s:e].strip()


def build():
    md = SOURCE.read_text(encoding="utf-8")
    zh_abs = clean_md(extract_block(md, "## 摘要", "## Abstract"))
    en_abs = clean_md(extract_block(md, "## Abstract", "## 1 引言"))
    zh_abs = re.sub(r"\n\n\*\*关键词：\*\*.*", "", zh_abs, flags=re.S).strip()
    en_abs = re.sub(r"\n\n\*\*Keywords:\*\*.*", "", en_abs, flags=re.S).strip()
    body = md[md.index("## 1 引言"):]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.55)
    sec.right_margin = Cm(1.55)
    sec.header_distance = Cm(0.65)
    sec.footer_distance = Cm(0.65)
    set_columns(sec, 1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("SOH estimation | Auditable reproduction and generalization boundaries")
    set_run_font(hr, size=7.2, italic=True, color=MUTED)
    add_page_field(sec.footer.paragraphs[0])

    journal = doc.add_paragraph()
    journal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    journal.paragraph_format.space_after = Pt(8)
    r = journal.add_run("Journal-ready research manuscript")
    set_run_font(r, size=8, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(TITLE)
    set_run_font(r, size=17.5, bold=True, east_asia="SimHei")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(EN_TITLE)
    set_run_font(r, size=10.5, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Anonymous author(s) - affiliations withheld for double-blind review")
    set_run_font(r, size=9.2, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run("Configuration hash: 797a026852b95eecdc1ce2a284d21396bd95df0df5ba66e533805076d8933a10")
    set_run_font(r, size=7.2, color=MUTED)

    info = doc.add_table(rows=2, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Inches(2.0)
    info.columns[1].width = Inches(4.95)
    labels = (("A R T I C L E   I N F O", "A B S T R A C T"),)
    for c, txt in enumerate(labels[0]):
        p = info.cell(0, c).paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(txt)
        set_run_font(r, size=8.5, bold=True)
    left = info.cell(1, 0)
    left.text = ""
    p = left.paragraphs[0]
    r = p.add_run("Keywords:\nLithium-ion battery\nState of health\nSmall-sample learning\nPhysical consistency\nDomain alignment\nNegative transfer")
    set_run_font(r, size=7.8, italic=True)
    right = info.cell(1, 1)
    right.text = ""
    add_inline(right.paragraphs[0], en_abs, size=7.7, justify=True)
    for row in info.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=60, start=100, bottom=60, end=100)
    set_table_borders(info, color="FFFFFF", size="0")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("中文摘要")
    set_run_font(r, size=9, bold=True)
    p = doc.add_paragraph()
    add_inline(p, zh_abs, size=8.2, justify=True)

    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    body_sec.page_width = sec.page_width
    body_sec.page_height = sec.page_height
    body_sec.top_margin = sec.top_margin
    body_sec.bottom_margin = sec.bottom_margin
    body_sec.left_margin = sec.left_margin
    body_sec.right_margin = sec.right_margin
    body_sec.header_distance = sec.header_distance
    body_sec.footer_distance = sec.footer_distance
    set_columns(body_sec, 2, space_twips=360)
    body_sec.header.is_linked_to_previous = True
    body_sec.footer.is_linked_to_previous = True

    lines = body.splitlines()
    i = 0
    figures_added = False
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped or stripped.startswith("<!--"):
            i += 1
            continue
        if stripped.startswith("|---"):
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, rows)
            continue
        if stripped.startswith("## "):
            heading = stripped[3:]
            if heading == "5 讨论" and not figures_added:
                add_figure(doc, "cila_outer_rmse_multiseed_log_v1.png", "Fig. 1. Frozen outer-test RMSE across five seeds (log scale).")
                add_figure(doc, "cila_outer_paired_delta_v1.png", "Fig. 2. Seed-wise paired RMSE differences; negative values favor CILA.")
                add_figure(doc, "cila_outer_prediction_iqr_v1.png", "Fig. 3. Outer-test prediction dispersion and recursive extrapolation failure.")
                add_figure(doc, "cila_outer_bootstrap_ci_v1.png", "Fig. 4. Circular time-block bootstrap 95% confidence intervals.")
                figures_added = True
            add_heading(doc, heading, 1)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 2)
            i += 1
            continue
        if stripped.startswith(">"):
            i += 1
            continue
        if re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style=None)
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.first_line_indent = Cm(-0.2)
            add_inline(p, "• " + stripped[2:], size=8.5, justify=False)
            i += 1
            continue

        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("<!--"):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.35)
        add_inline(p, " ".join(para_lines), size=8.7, justify=True)

    # Keep the article-like reference density readable and balance the final two-column page.
    reference_index = 0
    for p in doc.paragraphs:
        if p.text.strip().startswith("[") and re.match(r"^\[\d+\]", p.text.strip()):
            reference_index += 1
            if reference_index == 9:
                breaker = p.insert_paragraph_before()
                breaker.paragraph_format.space_after = Pt(0)
                breaker.add_run().add_break(WD_BREAK.COLUMN)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                set_run_font(run, size=7.0, color=BLUE)

    props = doc.core_properties
    props.title = TITLE
    props.subject = "Lithium-ion battery SOH estimation, auditable reproduction, physics constraints, and domain alignment"
    props.author = "Anonymous author(s)"
    props.keywords = "SOH, lithium-ion battery, TimesNet, iTransformer, physical constraint, domain alignment"

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
